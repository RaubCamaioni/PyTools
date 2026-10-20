# xlsx, excel, csv, json, spreadsheet, converter

from pathlib import Path
from typing import Literal
import csv
import json
import zipfile
from openpyxl import load_workbook


def xlsx_to_csv(
    spreadsheet: Path,
    output: Literal["csv", "tsv", "json", "markdown"] = "csv",
    sheet: str = "all",
) -> Path:
    """Convert an Excel .xlsx workbook to CSV, TSV, JSON (list of row objects) or a Markdown table. One sheet gives one file; all sheets give a zip (or one JSON object keyed by sheet)."""
    wb = load_workbook(str(spreadsheet), read_only=True, data_only=True)
    names = wb.sheetnames if sheet.strip().lower() == "all" else [sheet]
    if sheet.strip().lower() != "all" and sheet not in wb.sheetnames:
        if sheet.isdigit() and 1 <= int(sheet) <= len(wb.sheetnames):
            names = [wb.sheetnames[int(sheet) - 1]]
        else:
            raise ValueError(f"No sheet named '{sheet}'. Sheets: {', '.join(wb.sheetnames)}")

    def rows_of(name: str) -> list[list]:
        ws = wb[name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            if row is None or all(v is None for v in row):
                continue
            rows.append(["" if v is None else v for v in row])
        return rows

    def render(rows: list[list]) -> str:
        if output == "json":
            if not rows:
                return "[]"
            header = [str(h) for h in rows[0]]
            return json.dumps([dict(zip(header, r)) for r in rows[1:]], indent=2, default=str, ensure_ascii=False)
        if output == "markdown":
            if not rows:
                return ""
            width = max(len(r) for r in rows)
            rows = [[str(c) for c in r] + [""] * (width - len(r)) for r in rows]
            lines = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * width]
            lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
            return "\n".join(lines) + "\n"
        import io

        buf = io.StringIO()
        csv.writer(buf, delimiter="\t" if output == "tsv" else ",", lineterminator="\n").writerows(rows)
        return buf.getvalue()

    ext = {"csv": "csv", "tsv": "tsv", "json": "json", "markdown": "md"}[output]
    if len(names) == 1:
        out = spreadsheet.with_suffix("." + ext)
        out.write_text(render(rows_of(names[0])), encoding="utf-8")
        return out
    if output == "json":
        out = spreadsheet.with_suffix(".json")
        data = {name: json.loads(render(rows_of(name))) for name in names}
        out.write_text(json.dumps(data, indent=2, ensure_ascii=False, default=str), encoding="utf-8")
        return out
    out = spreadsheet.with_name(f"{spreadsheet.stem}_{ext}.zip")
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
        for name in names:
            safe = "".join(c if c.isalnum() or c in "-_ " else "_" for c in name)
            zf.writestr(f"{safe}.{ext}", render(rows_of(name)))
    return out
