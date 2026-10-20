# csv, xlsx, excel, json, spreadsheet, converter

from pathlib import Path
from typing import Literal
import csv
import json
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


def csv_to_xlsx(
    data_file: Path,
    input_format: Literal["csv_or_tsv", "json"] = "csv_or_tsv",
    numbers: Literal["convert", "keep_as_text"] = "convert",
) -> Path:
    """Convert a CSV/TSV file or a JSON list of objects into an Excel .xlsx workbook with a bold header row, auto-sized columns and real numbers."""
    text = data_file.read_text(encoding="utf-8-sig", errors="replace")

    if input_format == "json":
        data = json.loads(text)
        if isinstance(data, dict):
            data = [data]
        if not isinstance(data, list) or not all(isinstance(r, dict) for r in data):
            raise ValueError("JSON must be a list of objects")
        header: list[str] = []
        for row in data:
            for k in row:
                if k not in header:
                    header.append(k)
        rows = [header] + [[(json.dumps(r.get(k)) if isinstance(r.get(k), (dict, list)) else r.get(k, "")) for k in header] for r in data]
    else:
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        rows = list(csv.reader(text.splitlines(), dialect=dialect))

    def coerce(value):
        if numbers == "keep_as_text" or not isinstance(value, str):
            return value
        s = value.strip()
        try:
            return int(s)
        except ValueError:
            pass
        try:
            return float(s)
        except ValueError:
            return value

    wb = Workbook()
    ws = wb.active
    ws.title = data_file.stem[:31] or "Sheet1"
    widths: dict[int, int] = {}
    for r, row in enumerate(rows, 1):
        for c, value in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=coerce(value) if r > 1 else value)
            if r == 1:
                cell.font = Font(bold=True)
            widths[c] = min(60, max(widths.get(c, 8), len(str(value)) + 2))
    for c, w in widths.items():
        ws.column_dimensions[get_column_letter(c)].width = w
    if rows:
        ws.freeze_panes = "A2"
    out = data_file.with_suffix(".xlsx")
    wb.save(str(out))
    return out
