# pdf, excel, xlsx, csv, table, converter

from pathlib import Path
from typing import Literal
import csv
import zipfile
import pymupdf
from openpyxl import Workbook
from openpyxl.styles import Font


def pdf_to_excel(pdf: Path, output: Literal["xlsx", "csv"] = "xlsx", pages: str = "all") -> dict:
    """Extract the tables from a PDF into an Excel workbook (one sheet per table) or CSV files (zipped). Works on PDFs with real text, not scans."""
    with pymupdf.open(pdf) as doc:
        if pages.strip().lower() == "all":
            indexes = range(doc.page_count)
        else:
            indexes = []
            for part in pages.replace(" ", "").split(","):
                a, _, b = part.partition("-")
                if a:
                    indexes.extend(range(int(a) - 1, int(b or a)))
            indexes = [i for i in indexes if 0 <= i < doc.page_count]

        tables: list[tuple[str, list[list]]] = []
        for i in indexes:
            page = doc[i]
            found = page.find_tables()
            for n, table in enumerate(found.tables, 1):
                rows = table.extract()
                rows = [[("" if c is None else str(c).replace("\n", " ")) for c in row] for row in rows if any(c not in (None, "") for c in row)]
                if rows:
                    tables.append((f"p{i + 1}_t{n}", rows))

    if not tables:
        return {"tables_found": 0, "message": "No tables detected. If the PDF is a scan, OCR it first; if the table has no ruling lines, detection may miss it."}

    if output == "xlsx":
        wb = Workbook()
        wb.remove(wb.active)
        for name, rows in tables:
            ws = wb.create_sheet(name[:31])
            for r, row in enumerate(rows, 1):
                for c, value in enumerate(row, 1):
                    cell = ws.cell(row=r, column=c, value=_number(value))
                    if r == 1:
                        cell.font = Font(bold=True)
        out = pdf.with_suffix(".xlsx")
        wb.save(str(out))
    else:
        out = pdf.with_name(f"{pdf.stem}_tables.zip")
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
            for name, rows in tables:
                import io

                buf = io.StringIO()
                csv.writer(buf, lineterminator="\n").writerows(rows)
                zf.writestr(f"{pdf.stem}_{name}.csv", buf.getvalue())
    return {"file": out, "tables_found": len(tables), "rows_per_table": {n: len(r) for n, r in tables}}


def _number(value: str):
    text = value.strip().replace(",", "")
    try:
        return int(text)
    except ValueError:
        pass
    try:
        return float(text)
    except ValueError:
        return value
