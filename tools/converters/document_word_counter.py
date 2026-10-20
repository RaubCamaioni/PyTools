# word count, docx, pdf, pptx, pages, document, statistics

from pathlib import Path
import re


def document_word_counter(document: Path, words_per_page: int = 300) -> dict:
    """Word, character, sentence and paragraph counts for a Word, PDF, PowerPoint,
    Markdown or text file, with reading time and an estimated page count.
    """
    suffix = document.suffix.lower()
    if suffix == ".docx":
        from docx import Document

        doc = Document(str(document))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            parts += [cell.text for row in table.rows for cell in row.cells]
        text = "\n".join(parts)
        pages = None
    elif suffix == ".pdf":
        import pymupdf

        with pymupdf.open(document) as pdf:
            pages = pdf.page_count
            text = "\n".join(page.get_text() for page in pdf)
    elif suffix == ".pptx":
        from pptx import Presentation

        prs = Presentation(str(document))
        text = "\n".join(shape.text for slide in prs.slides for shape in slide.shapes if shape.has_text_frame)
        pages = len(prs.slides)
    else:
        text = document.read_text(encoding="utf-8-sig", errors="replace")
        pages = None
    words = re.findall(r"[^\s]+", text)
    letters = re.findall(r"[A-Za-z]", text)
    sentences = [s for s in re.split(r"[.!?]+(?:\s|$)", text) if s.strip()]
    paragraphs = [p for p in re.split(r"\n\s*\n", text) if p.strip()]
    return {
        "file": document.name,
        "words": len(words),
        "characters": len(text),
        "characters_without_spaces": len(re.sub(r"\s", "", text)),
        "letters": len(letters),
        "sentences": len(sentences),
        "paragraphs": len(paragraphs),
        "pages_in_file": pages,
        "estimated_pages_at_words_per_page": round(len(words) / max(1, words_per_page), 1),
        "reading_time_minutes": round(len(words) / 200, 1),
        "speaking_time_minutes": round(len(words) / 130, 1),
    }
