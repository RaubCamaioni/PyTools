# excel, xlsx, csv, html, xml, table, converter

from pathlib import Path
from typing import Literal
import csv
import html
import re
from openpyxl import load_workbook


def excel_to_html(spreadsheet: Path, output: Literal["html", "xml"] = "html") -> Path:
    """Turn a spreadsheet (xlsx or csv) into an HTML table page or a simple XML file
    with one element per row and one child per column.
    """
    sheets = []
    if spreadsheet.suffix.lower() in (".csv", ".tsv", ".txt"):
        text = spreadsheet.read_text(encoding="utf-8-sig", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        sheets.append((spreadsheet.stem, [list(r) for r in csv.reader(text.splitlines(), dialect=dialect)]))
    else:
        wb = load_workbook(str(spreadsheet), read_only=True, data_only=True)
        for ws in wb.worksheets:
            rows = [
                ["" if v is None else v for v in r]
                for r in ws.iter_rows(values_only=True)
                if any(v is not None for v in r)
            ]
            if rows:
                sheets.append((ws.title, rows))
    if not sheets:
        raise ValueError("No data found")

    def tag(name: str) -> str:
        clean = re.sub(r"[^A-Za-z0-9_]", "_", str(name)).strip("_") or "column"
        return clean if clean[0].isalpha() or clean[0] == "_" else "_" + clean

    if output == "xml":
        parts = ['<?xml version="1.0" encoding="utf-8"?>', "<workbook>"]
        for title, rows in sheets:
            parts.append(f'  <sheet name="{html.escape(str(title))}">')
            header = [tag(h) for h in rows[0]]
            for r in rows[1:]:
                parts.append("    <row>")
                for name, value in zip(header, r):
                    parts.append(f"      <{name}>{html.escape(str(value))}</{name}>")
                parts.append("    </row>")
            parts.append("  </sheet>")
        parts.append("</workbook>")
        out = spreadsheet.with_suffix(".xml")
        out.write_text("\n".join(parts) + "\n", encoding="utf-8")
        return out

    body = []
    for title, rows in sheets:
        body.append(f"<h2>{html.escape(str(title))}</h2><table>")
        body.append(
            "<thead><tr>" + "".join(f"<th>{html.escape(str(c))}</th>" for c in rows[0]) + "</tr></thead><tbody>"
        )
        for r in rows[1:]:
            body.append("<tr>" + "".join(f"<td>{html.escape(str(c))}</td>" for c in r) + "</tr>")
        body.append("</tbody></table>")
    style = (
        "<style>body{font-family:sans-serif}table{border-collapse:collapse;margin-bottom:2em}"
        "td,th{border:1px solid #ccc;padding:4px 8px}th{background:#eee}</style>"
    )
    out = spreadsheet.with_suffix(".html")
    out.write_text(
        f"<!doctype html><html><head><meta charset='utf-8'>{style}</head><body>{''.join(body)}</body></html>\n",
        encoding="utf-8",
    )
    return out
