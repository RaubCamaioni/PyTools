# pptx, powerpoint, text, markdown, slides, converter

from pathlib import Path
from typing import Literal
from pptx import Presentation


def pptx_to_text(pptx_file: Path, output: Literal["markdown", "plain_text"] = "markdown", include_notes: Literal["yes", "no"] = "yes") -> Path:
    """Extract all text from a PowerPoint .pptx (titles, bullets, tables and speaker notes), one section per slide, as Markdown or plain text."""
    prs = Presentation(str(pptx_file))
    md = output == "markdown"
    lines: list[str] = []
    for number, slide in enumerate(prs.slides, 1):
        title = slide.shapes.title.text.strip() if slide.shapes.title is not None and slide.shapes.title.has_text_frame else ""
        lines.append(f"## Slide {number}: {title}" if md else f"=== Slide {number}: {title} ===")
        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in para.runs).strip()
                    if text:
                        indent = "  " * para.level
                        lines.append(f"{indent}- {text}" if md else f"{indent}{text}")
            elif getattr(shape, "has_table", False) and shape.has_table:
                rows = [[cell.text.strip().replace("\n", " ") for cell in row.cells] for row in shape.table.rows]
                if md and rows:
                    lines.append("| " + " | ".join(rows[0]) + " |")
                    lines.append("|" + "---|" * len(rows[0]))
                    lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
                else:
                    lines += ["\t".join(r) for r in rows]
        if include_notes == "yes" and slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(("> **Notes:** " if md else "Notes: ") + notes.replace("\n", " "))
        lines.append("")
    out = pptx_file.with_suffix(".md" if md else ".txt")
    out.write_text("\n".join(lines).strip() + "\n", encoding="utf-8")
    return out
