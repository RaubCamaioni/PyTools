# docx, word, merge, combine, join

from pathlib import Path
from typing import Literal
import copy
from docx import Document
from docx.enum.text import WD_BREAK


def docx_merge(
    first_docx: Path,
    second_docx: Path,
    third_docx: Path,
    separator: Literal["page_break", "blank_line", "nothing"] = "page_break",
) -> dict:
    """Join two or three Word documents into one, in order, with a page break (or a
    blank line, or nothing) between them. Upload the same file twice if you only have
    two.
    """
    files, seen = [], set()
    for f in (first_docx, second_docx, third_docx):
        key = (f.name, f.stat().st_size)
        if key not in seen:
            seen.add(key)
            files.append(f)
    base = Document(str(files[0]))
    for extra in files[1:]:
        if separator == "page_break":
            base.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif separator == "blank_line":
            base.add_paragraph()
        other = Document(str(extra))
        for element in other.element.body:
            if element.tag.endswith("}sectPr"):
                continue
            base.element.body.append(copy.deepcopy(element))
    out = Path(f"{files[0].stem}_merged.docx")
    base.save(str(out))
    return {"file": out, "documents": len(files), "paragraphs": len(base.paragraphs)}
