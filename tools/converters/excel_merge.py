# excel, xlsx, csv, merge, combine, sheets

from pathlib import Path
from typing import Literal
import csv
from openpyxl import Workbook, load_workbook


def excel_merge(
    first_file: Path,
    second_file: Path,
    third_file: Path,
    mode: Literal["one_sheet_per_file", "stack_rows_single_header"] = "one_sheet_per_file",
) -> dict:
    """Combine two or three spreadsheets (xlsx or csv) into one workbook: each file as
    its own sheet, or all rows stacked under one header. Upload the same file twice
    if you have two.
    """
    files, seen = [], set()
    for f in (first_file, second_file, third_file):
        key = (f.name, f.stat().st_size)
        if key not in seen:
            seen.add(key)
            files.append(f)

    def rows_of(path: Path) -> list[list]:
        if path.suffix.lower() in (".csv", ".tsv", ".txt"):
            text = path.read_text(encoding="utf-8-sig", errors="replace")
            try:
                dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
            except csv.Error:
                dialect = csv.excel
            return [list(r) for r in csv.reader(text.splitlines(), dialect=dialect)]
        wb = load_workbook(str(path), read_only=True, data_only=True)
        return [list(r) for r in wb.worksheets[0].iter_rows(values_only=True) if any(v is not None for v in r)]

    out_wb = Workbook()
    out_wb.remove(out_wb.active)
    total = 0
    if mode == "one_sheet_per_file":
        for f in files:
            ws = out_wb.create_sheet(f.stem[:31])
            for r in rows_of(f):
                ws.append(r)
                total += 1
    else:
        ws = out_wb.create_sheet("merged")
        header_done = False
        for f in files:
            rows = rows_of(f)
            if not rows:
                continue
            if not header_done:
                ws.append(rows[0])
                header_done = True
            for r in rows[1:]:
                ws.append(r)
                total += 1
    out = Path(f"{files[0].stem}_merged.xlsx")
    out_wb.save(str(out))
    return {"file": out, "files": len(files), "rows": total}
