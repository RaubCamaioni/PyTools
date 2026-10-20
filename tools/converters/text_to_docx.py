# text, docx, word, markdown, document, converter

from pathlib import Path
from typing import Literal
import re
from docx import Document
from docx.shared import Pt


def add_inline(paragraph, text: str):
    """Add text with **bold**, *italic* and `code` runs."""
    for token in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text):
        if not token:
            continue
        if token.startswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
        elif token.startswith("*"):
            paragraph.add_run(token[1:-1]).italic = True
        else:
            paragraph.add_run(token)


def text_to_docx(text_file: Path, input_format: Literal["markdown", "plain_text"] = "markdown", font_size: int = 11) -> Path:
    """Create a Word .docx from a Markdown or plain text file (headings, bullet and numbered lists, bold/italic, code blocks, simple tables)."""
    text = text_file.read_text(encoding="utf-8-sig", errors="replace")
    doc = Document()
    doc.styles["Normal"].font.size = Pt(max(6, min(font_size, 24)))

    if input_format == "plain_text":
        for para in re.split(r"\n\s*\n", text):
            doc.add_paragraph(para.strip())
        out = text_file.with_suffix(".docx")
        doc.save(str(out))
        return out

    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            code = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                code.append(lines[i])
                i += 1
            p = doc.add_paragraph()
            run = p.add_run("\n".join(code))
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        elif m := re.match(r"^(#{1,6})\s+(.*)", line):
            doc.add_heading(m.group(2).strip(), level=len(m.group(1)))
        elif m := re.match(r"^\s*[-*+]\s+(.*)", line):
            add_inline(doc.add_paragraph(style="List Bullet"), m.group(1))
        elif m := re.match(r"^\s*\d+[.)]\s+(.*)", line):
            add_inline(doc.add_paragraph(style="List Number"), m.group(1))
        elif line.strip().startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                    rows.append(cells)
                i += 1
            if rows:
                table = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
                table.style = "Table Grid"
                for r, row in enumerate(rows):
                    for c, cell in enumerate(row):
                        add_inline(table.cell(r, c).paragraphs[0], cell)
            continue
        elif m := re.match(r"^>\s?(.*)", line):
            p = doc.add_paragraph()
            add_inline(p, m.group(1))
            p.paragraph_format.left_indent = Pt(24)
        elif line.strip() in ("---", "***", "___"):
            doc.add_paragraph("_" * 40)
        elif line.strip():
            para = [line.strip()]
            while i + 1 < len(lines) and lines[i + 1].strip() and not re.match(r"^(#|\s*[-*+]\s|\s*\d+[.)]\s|```|\||>)", lines[i + 1]):
                i += 1
                para.append(lines[i].strip())
            add_inline(doc.add_paragraph(), " ".join(para))
        i += 1

    out = text_file.with_suffix(".docx")
    doc.save(str(out))
    return out
