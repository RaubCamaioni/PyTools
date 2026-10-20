# markdown, pptx, powerpoint, slides, converter

from pathlib import Path
import re
from pptx import Presentation
from pptx.util import Pt


def markdown_to_pptx(markdown_file: Path, title: str = "") -> dict:
    """Turn a Markdown file into a PowerPoint deck: each # or ## heading starts a slide,
    the bullets under it become the slide's body (nested bullets indent), paragraphs
    become plain lines.
    """
    text = markdown_file.read_text(encoding="utf-8-sig", errors="replace")
    slides: list[tuple[str, list[tuple[int, str]]]] = []
    current = None
    for line in text.splitlines():
        heading = re.match(r"^(#{1,2})\s+(.*)", line)
        if heading:
            current = (heading.group(2).strip(), [])
            slides.append(current)
            continue
        if current is None:
            if line.strip():
                current = (title or markdown_file.stem, [])
                slides.append(current)
            else:
                continue
        bullet = re.match(r"^(\s*)[-*+]\s+(.*)", line)
        numbered = re.match(r"^(\s*)\d+[.)]\s+(.*)", line)
        if bullet or numbered:
            m = bullet or numbered
            level = min(len(m.group(1)) // 2, 4)
            current[1].append((level, m.group(2).strip()))
        elif line.strip() and not line.startswith("```"):
            current[1].append((0, re.sub(r"[*_`]", "", line.strip())))
    if not slides:
        return {"error": "No headings or text found"}

    prs = Presentation()
    for heading, items in slides:
        layout = prs.slide_layouts[1] if items else prs.slide_layouts[5]
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = re.sub(r"[*_`]", "", heading)
        if items:
            body = slide.placeholders[1].text_frame
            body.clear()
            for i, (level, item) in enumerate(items[:20]):
                para = body.paragraphs[0] if i == 0 else body.add_paragraph()
                para.text = re.sub(r"[*_`]", "", item)
                para.level = level
                para.font.size = Pt(max(14, 24 - 3 * level))
    out = markdown_file.with_suffix(".pptx")
    prs.save(str(out))
    return {"file": out, "slides": len(slides)}
